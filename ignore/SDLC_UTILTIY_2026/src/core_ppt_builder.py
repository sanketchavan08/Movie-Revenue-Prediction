"""Generic PowerPoint helpers. Nothing in here knows about specific slides."""
import copy
import logging
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

log = logging.getLogger(__name__)

TOKEN_PATTERN = re.compile(r"\{\{[A-Z0-9_]+\}\}")


def _iter_shapes(shapes):
    """All shapes, including the ones inside grouped shapes."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _replace_in_paragraph(paragraph, token, value):
    """Replace a token even when PowerPoint has split it across several runs.

    Only the runs that contain the token are touched, so the rest of the
    paragraph keeps its own fonts/colours. The replacement takes the styling
    of the run where the token starts.
    """
    count = 0
    while True:
        runs = paragraph.runs
        full = "".join(r.text for r in runs)
        start = full.find(token)
        if start < 0:
            return count
        end = start + len(token)
        pos, first = 0, True
        for run in runs:
            run_start, run_end = pos, pos + len(run.text)
            pos = run_end
            if run_end <= start or run_start >= end:
                continue
            a, b = max(start, run_start) - run_start, min(end, run_end) - run_start
            if first:
                run.text = run.text[:a] + value + run.text[b:]
                first = False
            else:
                run.text = run.text[:a] + run.text[b:]
        count += 1


def _text_frames(shape):
    """Every text frame in a shape: normal text boxes and every table cell."""
    if shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


class PPTBuilder:
    def __init__(self, template_path):
        self.prs = Presentation(template_path)
        log.info("Loaded template: %s (%d slides)", template_path.name, len(self.prs.slides))

    # ------------------------------------------------------------ finding shapes
    def find_shape(self, name):
        """Find a shape by its Selection Pane name, anywhere in the deck."""
        hits = [(i, s) for i, slide in enumerate(self.prs.slides, start=1)
                for s in _iter_shapes(slide.shapes) if s.name == name]
        if len(hits) == 1:
            return hits[0][1]
        if not hits:
            names = {i: [s.name for s in _iter_shapes(slide.shapes)]
                     for i, slide in enumerate(self.prs.slides, start=1)}
            listing = "\n".join(f"  slide {i}: {n}" for i, n in names.items())
            raise ValueError(f"No shape named '{name}' in the template. "
                           f"Rename it in PowerPoint's Selection Pane. Shapes found:\n{listing}")
        raise ValueError(f"Shape name '{name}' is used {len(hits)} times "
                       f"(slides {[i for i, _ in hits]}); names must be unique.")

    # ------------------------------------------------------------ text
    def replace_tokens(self, values):
        """values: {'{{MON}}': 'Sep', ...}. Searches text boxes, tables and groups on every slide."""
        counts = dict.fromkeys(values, 0)
        for slide in self.prs.slides:
            for shape in _iter_shapes(slide.shapes):
                for tf in _text_frames(shape):
                    for p in tf.paragraphs:
                        for token, value in values.items():
                            counts[token] += _replace_in_paragraph(p, token, str(value))
        for token, n in counts.items():
            if n:
                log.info("Token %-16s -> '%s' (%d place%s)", token, values[token], n, "" if n == 1 else "s")
            else:
                log.warning("Token %s not found in the template", token)
        return counts

    def leftover_tokens(self):
        """Any {{TOKEN}} still in the deck after all replacements -> list of (slide_no, token)."""
        found = []
        for i, slide in enumerate(self.prs.slides, start=1):
            for shape in _iter_shapes(slide.shapes):
                for tf in _text_frames(shape):
                    for p in tf.paragraphs:
                        text = "".join(r.text for r in p.runs)
                        found += [(i, t) for t in TOKEN_PATTERN.findall(text)]
        return found

    # ------------------------------------------------------------ tables
    @staticmethod
    def set_cell_text(cell, text):
        """Replace a table cell's text but keep the font/size/colour of its first run."""
        tf = cell.text_frame
        p = tf.paragraphs[0]
        for extra in tf.paragraphs[1:]:                     # drop extra paragraphs
            extra._p.getparent().remove(extra._p)
        text = str(text)
        if text == "":
            # Empty cell: keep the font size for the empty line but remove the runs
            # (an empty run can be drawn at the default 18 pt and make the row tall).
            if p.runs:
                r_pr = p.runs[0]._r.find(qn("a:rPr"))
                if r_pr is not None and r_pr.get("sz"):
                    p._p.get_or_add_endParaRPr().set("sz", r_pr.get("sz"))
            for run in p.runs:
                run._r.getparent().remove(run._r)
        elif p.runs:
            p.runs[0].text = text
            for run in p.runs[1:]:
                run._r.getparent().remove(run._r)
        else:
            # Cell was empty in the template: give the new text the cell's own font settings.
            run = p.add_run()
            run.text = text
            end = p._p.find(qn("a:endParaRPr"))
            if end is not None:
                r_pr = run._r.get_or_add_rPr()
                for key, value in end.attrib.items():
                    r_pr.set(key, value)
                for child in end:
                    r_pr.append(copy.deepcopy(child))

    @staticmethod
    def set_cell_fill(cell, hex_color):
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(hex_color)

    @staticmethod
    def get_table(shape):
        if not shape.has_table:
            raise TypeError(f"Shape '{shape.name}' is not a table")
        return shape.table

    def fill_repeating_table(self, shape, template_row, rows, fills=None):
        """Rebuild a table with one row per item, styled like `template_row`.

        rows:  list of lists of cell texts, e.g. [['109360-1', 'db-COOP', '100.00', ...], ...]
        fills: optional {(row_index, col_index): 'RRGGBB'} for coloured cells (row_index 0 = first data row)
        Everything above template_row (the header) is kept; the old data rows are removed.
        """
        table = self.get_table(shape)
        tbl = table._tbl
        tr_list = tbl.tr_lst
        if template_row >= len(tr_list):
            raise ValueError(f"'{shape.name}' has no row {template_row} to use as the row template")
        pattern = copy.deepcopy(tr_list[template_row])
        for tr in tr_list[template_row:]:                    # remove old data rows
            tbl.remove(tr)

        if not rows:
            log.warning("'%s': no rows to show - leaving one empty row", shape.name)
            rows = [[""] * len(table.columns)]
        for r, values in enumerate(rows):
            tbl.append(copy.deepcopy(pattern))
            cells = table.rows[template_row + r].cells
            for c, text in enumerate(values):
                if c < len(cells):
                    self.set_cell_text(cells[c], text)
            for (fr, fc), color in (fills or {}).items():
                if fr == r:
                    self.set_cell_fill(cells[fc], color)

        shape.height = sum(row.height for row in table.rows)  # PowerPoint grows rows for wrapped text itself
        log.info("Table '%s': %d data rows", shape.name, len(rows))

    def fit_table(self, shape, template_row, max_bottom_in, min_font_pt):
        """Shrink the data rows' font so the table ends above `max_bottom_in` (inches from slide top).

        Row heights stored in the file are only minimums - PowerPoint grows a row to fit its text -
        so the drawn height is estimated from font size, cell margins and text wrapping.
        Does nothing when the table already fits. Warns if even `min_font_pt` is not enough.
        """
        table = self.get_table(shape)
        all_rows = list(table.rows)
        header_h = sum(r.height for r in all_rows[:template_row])
        data_rows = all_rows[template_row:]
        col_widths = [c.width for c in table.columns]
        available = Inches(max_bottom_in) - shape.top - header_h
        current_pt = next((run.font.size.pt for c in data_rows[0].cells for p in c.text_frame.paragraphs
                           for run in p.runs if run.font.size), 18.0)   # 18 pt = PowerPoint default
        margin = Inches(0.02)

        def drawn_height(font_pt, cell_margin, keep_row_heights=False):
            """Estimated drawn height of the data rows. keep_row_heights: use the template's row
            heights as a minimum (true for the current table, false once we resize the rows)."""
            total = 0
            for row in data_rows:
                lines = 1
                for cell, width in zip(row.cells, col_widths):
                    usable_pt = Emu(width - cell.margin_left - cell.margin_right).pt
                    chars_per_line = max(1, int(usable_pt / (font_pt * 0.55)))
                    lines = max(lines, -(-len(cell.text) // chars_per_line))   # ceiling division
                needed = Pt(lines * font_pt * 1.2) + 2 * cell_margin
                total += max(row.height, needed) if keep_row_heights else needed
            return total

        margins_now = data_rows[0].cells[0].margin_top + data_rows[0].cells[0].margin_bottom
        if drawn_height(current_pt, margins_now // 2, keep_row_heights=True) <= available:
            return

        font_pt = current_pt
        while font_pt > min_font_pt and drawn_height(font_pt, margin) > available:
            font_pt -= 0.5
        line_h = Pt(font_pt * 1.2) + 2 * margin
        for row in data_rows:
            row.height = line_h                           # minimum; wrapped rows grow by themselves
            for cell in row.cells:
                cell.margin_top = cell.margin_bottom = margin
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(font_pt)
                    # empty cells keep the default 18 pt line height unless this is set too
                    p._p.get_or_add_endParaRPr().set("sz", str(int(font_pt * 100)))
        shape.height = sum(r.height for r in table.rows)

        if drawn_height(font_pt, margin) > available:
            log.warning("'%s': %d rows at %s pt (the minimum) may not fit - check the slide",
                        shape.name, len(data_rows), font_pt)
        else:
            log.info("'%s': %d rows - tightened cell margins%s so the table fits", shape.name,
                     len(data_rows), "" if font_pt == current_pt else f" and font {current_pt} -> {font_pt} pt")

    # ------------------------------------------------------------ indicators
    @staticmethod
    def set_visible(shape, visible):
        c_nv_pr = shape._element.find(".//" + qn("p:cNvPr"))
        if visible:
            c_nv_pr.attrib.pop("hidden", None)
        else:
            c_nv_pr.set("hidden", "1")

    def set_indicator(self, shape, status, colors):
        """status 'green' -> green triangle pointing up; 'red' -> red pointing down."""
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            log.warning("'%s' is not a drawn shape (it is %s); cannot recolour it", shape.name, shape.shape_type)
            return
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(colors[status])
        shape.line.fill.background()            # no outline
        shape.rotation = 0 if status == "green" else 180
        self.set_visible(shape, True)

    # ------------------------------------------------------------ charts
    @staticmethod
    def update_line_chart(shape, categories, series, number_format, two_level=True):
        """Replace the chart's data; styling of the template chart is kept.

        categories: list of (year, month_label), e.g. [(2025, 'Sep'), ...]
        series:     list of (name, [values])  - None values are left as gaps
        """
        if not shape.has_chart:
            raise TypeError(f"Shape '{shape.name}' is not a chart")
        data = CategoryChartData(number_format=number_format)
        if two_level:                       # month labels with the year underneath
            parents = {}
            for year, label in categories:
                if year not in parents:
                    parents[year] = data.add_category(str(year))
                parents[year].add_sub_category(label)
        else:
            data.categories = [f"{label}-{str(year)[2:]}" for year, label in categories]
        for name, values in series:
            data.add_series(name, values)
        shape.chart.replace_data(data)

    # ------------------------------------------------------------ save
    def save(self, output_path):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.prs.save(output_path)
        except PermissionError:
            raise PermissionError(f"Cannot write {output_path.name} - close it in PowerPoint and run again.")
        log.info("Deck saved: %s", output_path)
