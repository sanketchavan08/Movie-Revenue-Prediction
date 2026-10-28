"""Reads the monthly raw Excel: single cells (KPI1 & KPI2 Summary) and list tabs (Applications View)."""
import logging
import re

import openpyxl
from openpyxl.utils import column_index_from_string

log = logging.getLogger(__name__)


def _norm(text):
    """Lower-case and collapse whitespace/newlines so header checks are forgiving."""
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _to_number(raw, field, spec):
    if raw is None or (isinstance(raw, str) and raw.strip() in ("", "-")):
        if spec.get("blank_as_zero"):
            return 0
        raise ValueError(f"'{field}' (cell {spec['cell']}) is blank, but a number is required")
    if isinstance(raw, (int, float)):
        return raw
    try:
        return float(str(raw).replace("%", "").replace(",", "").strip())
    except ValueError:
        raise ValueError(f"'{field}' (cell {spec['cell']}) is not a number: '{raw}'")


def read_summary(excel_path, sheet_cfg):
    """Returns {field_name: value} for every field listed in mapping.json -> summary_sheet."""
    sheet_name = sheet_cfg["sheet"]
    log.info("Reading '%s' from %s", sheet_name, excel_path.name)

    # read_only=True opens a big multi-sheet workbook much faster.
    # data_only=True returns calculated values instead of formulas.
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Sheets in file: {wb.sheetnames}")
        ws = wb[sheet_name]

        values = {}
        for field, spec in sheet_cfg["fields"].items():
            # Safety check: the header above the value must still say what we expect.
            # If the report layout ever shifts, we stop instead of putting a wrong number on a slide.
            if spec.get("header_cell"):
                actual = ws[spec["header_cell"]].value
                if _norm(spec["header_contains"]) not in _norm(actual):
                    raise ValueError(
                        f"Layout changed? Expected header containing '{spec['header_contains']}' "
                        f"in {spec['header_cell']}, found '{actual}'")

            raw = ws[spec["cell"]].value   # merged cells keep their value in the top-left cell
            if spec.get("type") == "text":
                values[field] = str(raw or "").strip()
            else:
                values[field] = _to_number(raw, field, spec)
            log.info("  %-24s %-4s = %s", field, spec["cell"], values[field])
        return values
    finally:
        wb.close()


def read_table(excel_path, table_cfg):
    """Reads a list-style tab (one row per application) into a list of dicts, in Excel order.

    Stops at the first row whose key column is blank. Every column's header is checked first.
    """
    sheet_name = table_cfg["sheet"]
    log.info("Reading '%s' from %s", sheet_name, excel_path.name)
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Sheets in file: {wb.sheetnames}")
        ws = wb[sheet_name]
        cols = {f: column_index_from_string(spec["col"]) for f, spec in table_cfg["columns"].items()}

        header = next(ws.iter_rows(min_row=table_cfg["header_row"], max_row=table_cfg["header_row"],
                                   values_only=True))
        for field, spec in table_cfg["columns"].items():
            actual = header[cols[field] - 1] if cols[field] <= len(header) else None
            if _norm(spec["header_contains"]) not in _norm(actual):
                raise ValueError(
                    f"Layout changed? '{sheet_name}' column {spec['col']} should contain "
                    f"'{spec['header_contains']}', found '{actual}'")

        key_col = cols[table_cfg["key_field"]]
        records = []
        for excel_row, values in enumerate(ws.iter_rows(min_row=table_cfg["first_data_row"], values_only=True),
                                           start=table_cfg["first_data_row"]):
            values = list(values) + [None] * (max(cols.values()) - len(values))
            if values[key_col - 1] in (None, ""):
                break                                        # end of the list
            record = {"_excel_row": excel_row}
            for field, spec in table_cfg["columns"].items():
                raw = values[cols[field] - 1]
                if spec.get("type") == "text":
                    record[field] = str(raw or "").strip()
                else:
                    record[field] = _to_number(raw, field, {**spec, "cell": f"{spec['col']}{excel_row}"})
            records.append(record)
        log.info("  %d rows read (Excel rows %d-%d)", len(records), table_cfg["first_data_row"],
                 table_cfg["first_data_row"] + len(records) - 1)
        return records
    finally:
        wb.close()


def read_grouped_values(excel_path, cfg):
    """Pivot-style read of a raw-data tab: filter rows, then collect distinct values per key.

    Columns are found by their exact header text, so a moved column does not break anything.
    Returns {key: [value, value, ...]}  e.g. {'131506-1': ['RMG0815606', 'RMG0815632']}
    """
    sheet_name = cfg["sheet"]
    log.info("Reading '%s' from %s", sheet_name, excel_path.name)
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Sheets in file: {wb.sheetnames}")
        ws = wb[sheet_name]
        header = next(ws.iter_rows(min_row=cfg["header_row"], max_row=cfg["header_row"], values_only=True))
        positions = {}
        for name in [cfg["key_header"], cfg["value_header"], *cfg["filters"]]:
            hits = [i for i, h in enumerate(header) if _norm(h) == _norm(name)]
            if len(hits) != 1:
                raise ValueError(f"'{sheet_name}' row {cfg['header_row']}: expected one column headed "
                                 f"'{name}', found {len(hits)}")
            positions[name] = hits[0]

        wanted = {name: _norm(value) for name, value in cfg["filters"].items()}
        key_i, value_i = positions[cfg["key_header"]], positions[cfg["value_header"]]
        grouped, matched, scanned = {}, 0, 0
        for row in ws.iter_rows(min_row=cfg["first_data_row"], values_only=True):
            if key_i >= len(row) or row[key_i] in (None, ""):
                continue                                   # blank row - skip, don't stop
            scanned += 1
            if all(_norm(row[positions[n]]) == v for n, v in wanted.items()):
                matched += 1
                values = grouped.setdefault(str(row[key_i]).strip(), [])
                value = str(row[value_i] or "").strip()
                if value and value not in values:
                    values.append(value)
        log.info("  %d rows scanned, %d match %s -> %d keys", scanned, matched,
                 " & ".join(f"{n}={v}" for n, v in cfg["filters"].items()), len(grouped))
        return grouped
    finally:
        wb.close()
