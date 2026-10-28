"""Reads and updates SDLC_PROGRESS_GRAPHS.xlsx (tab KPI_PROGRESS).

Layout: Year (merged per year) | Month | Release KPI | Automation KPI
Values are stored as fractions (1.0 = 100%, 0.956 = 95.60%).

NOTE: openpyxl does not keep charts when it saves. Any chart inside this
workbook is removed on save - the deck's chart replaces it. A timestamped
backup is written before every save.
"""
import logging
import shutil
from copy import copy
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string

from utils_config import MONTH_ABBR

log = logging.getLogger(__name__)


def _month_number(value):
    """'Sep' / 'September' / 'SEP' / a real date -> 9.  None if not a month."""
    if isinstance(value, (datetime, date)):
        return value.month
    text = str(value or "").strip()[:3].title()
    return MONTH_ABBR.index(text) + 1 if text in MONTH_ABBR else None


def _to_fraction(value):
    """1 / 0.956 / '95.6%' / 95.6  ->  fraction (0.956). Blank -> None."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str):
        text = value.strip()
        number = float(text.rstrip("%"))
        return number / 100 if text.endswith("%") or number > 1.5 else number
    number = float(value)
    return number / 100 if number > 1.5 else number


class HistoryBook:
    def __init__(self, path, layout):
        self.path = path
        self.wb = openpyxl.load_workbook(path)   # normal mode: we will write to it
        if layout["sheet"] not in self.wb.sheetnames:
            raise ValueError(f"Sheet '{layout['sheet']}' not in {path.name}: {self.wb.sheetnames}")
        self.ws = self.wb[layout["sheet"]]
        self.first_row = layout["first_data_row"]
        self.col = {k: column_index_from_string(layout[f"{k}_col"])
                    for k in ("year", "month", "release", "automation")}
        self.rows = {}        # date(year, month, 1) -> row number
        self.last_row = self.first_row - 1
        self._warned = set()
        self._index_rows()

    # ------------------------------------------------------------ reading
    def _index_rows(self):
        ws, ycol = self.ws, self.col["year"]

        # The Year column is merged per year, so only the top cell holds the value.
        year_of_row = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_col <= ycol <= rng.max_col:
                top_value = ws.cell(rng.min_row, ycol).value
                for r in range(rng.min_row, rng.max_row + 1):
                    year_of_row[r] = top_value

        current_year, previous = None, None
        for r in range(self.first_row, ws.max_row + 1):
            month = _month_number(ws.cell(r, self.col["month"]).value)
            if month is None:
                continue
            year_cell = year_of_row.get(r, ws.cell(r, ycol).value)
            if year_cell not in (None, ""):
                current_year = int(year_cell)
            if current_year is None:
                raise ValueError(f"{self.path.name}: no Year found for row {r}")

            key = date(current_year, month, 1)
            if previous and (previous.year * 12 + previous.month + 1) != (key.year * 12 + key.month):
                log.warning("History rows are not in month order at row %s (%s after %s)",
                            r, f"{key:%b %Y}", f"{previous:%b %Y}")
            self.rows[key] = r
            self.last_row = r
            previous = key

    def get(self, month):
        """-> (release_fraction, automation_fraction); None for anything missing."""
        r = self.rows.get(month)
        if r is None:
            return None, None
        return (_to_fraction(self.ws.cell(r, self.col["release"]).value),
                _to_fraction(self.ws.cell(r, self.col["automation"]).value))

    def window(self, months):
        """[{'month': date, 'release': 0.99, 'automation': 0.95}, ...] for the given months."""
        out = []
        for m in months:
            release, automation = self.get(m)
            if (release is None or automation is None) and m not in self._warned:
                self._warned.add(m)
                log.warning("History has no complete values for %s", f"{m:%b %Y}")
            out.append({"month": m, "release": release, "automation": automation})
        return out

    # ------------------------------------------------------------ writing
    def upsert(self, month, release_frac, automation_frac):
        """Write this month's values. Re-running the same month overwrites, never duplicates."""
        r = self.rows.get(month)
        if r is None:
            r = self._append_row(month)
            log.info("History: added new row %s for %s", r, f"{month:%b %Y}")

        for key, new in (("release", release_frac), ("automation", automation_frac)):
            cell = self.ws.cell(r, self.col[key])
            old = _to_fraction(cell.value)
            if old is not None and abs(old - new) > 1e-6:
                log.warning("History %s %s: overwriting %.4f with %.4f", f"{month:%b %Y}", key, old, new)
            cell.value = new
            # Use the same % format as the row above (e.g. '0%' for Release, '0.00%' for Automation)
            above = self.ws.cell(r - 1, self.col[key]).number_format or ""
            cell.number_format = above if "%" in above else "0.00%"
        log.info("History %s -> release %.4f, automation %.4f", f"{month:%b %Y}", release_frac, automation_frac)

    def _append_row(self, month):
        if self.rows:
            latest = max(self.rows)
            expected = date(latest.year + latest.month // 12, latest.month % 12 + 1, 1)
            if month != expected:
                raise ValueError(f"History ends at {latest:%b %Y}; cannot jump to {month:%b %Y}. "
                                 f"Add the missing month rows first.")
        r = self.last_row + 1
        for c in self.col.values():                     # copy borders/fonts/number formats
            src, dst = self.ws.cell(r - 1, c), self.ws.cell(r, c)
            if src.has_style:
                dst._style = copy(src._style)
        self.ws.cell(r, self.col["year"]).value = month.year
        self.ws.cell(r, self.col["month"]).value = MONTH_ABBR[month.month - 1]
        self.rows[month] = r
        self.last_row = r
        return r

    def save(self, backup_dir):
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / f"{self.path.stem}_{datetime.now():%Y%m%d_%H%M%S}{self.path.suffix}"
        shutil.copy2(self.path, backup)
        try:
            self.wb.save(self.path)
        except PermissionError:
            raise PermissionError(f"Cannot write {self.path.name} - close it in Excel and run again.")
        log.info("History saved (backup: %s)", backup.name)
