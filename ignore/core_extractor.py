"""Reads the monthly raw Excel (KPI1 & KPI2 Summary tab) into a plain dict."""
import logging
import re

import openpyxl

log = logging.getLogger(__name__)


def _norm(text):
    """Lower-case and collapse whitespace/newlines so header checks are forgiving."""
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _to_number(raw, field, spec):
    if raw is None or (isinstance(raw, str) and raw.strip() in ("", "-")):
        if spec.get("blank_as_zero"):
            return 0
        raise ValueError(f"'{field}' (cell {spec['cell']}) is blank, but a number is required")
    if isinstance(raw, (int, float)):
        return raw
    try:
        return float(str(raw).replace("%", "").replace(",", "").strip())
    except ValueError:
        raise ValueError(f"'{field}' (cell {spec['cell']}) is not a number: '{raw}'")


def read_summary(excel_path, sheet_cfg):
    """Returns {field_name: value} for every field listed in mapping.json -> summary_sheet."""
    sheet_name = sheet_cfg["sheet"]
    log.info("Reading '%s' from %s", sheet_name, excel_path.name)

    # read_only=True opens a big multi-sheet workbook much faster.
    # data_only=True returns calculated values instead of formulas.
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Sheets in file: {wb.sheetnames}")
        ws = wb[sheet_name]

        values = {}
        for field, spec in sheet_cfg["fields"].items():
            # Safety check: the header above the value must still say what we expect.
            # If the report layout ever shifts, we stop instead of putting a wrong number on a slide.
            if spec.get("header_cell"):
                actual = ws[spec["header_cell"]].value
                if _norm(spec["header_contains"]) not in _norm(actual):
                    raise ValueError(
                        f"Layout changed? Expected header containing '{spec['header_contains']}' "
                        f"in {spec['header_cell']}, found '{actual}'")

            raw = ws[spec["cell"]].value   # merged cells keep their value in the top-left cell
            if spec.get("type") == "text":
                values[field] = str(raw or "").strip()
            else:
                values[field] = _to_number(raw, field, spec)
            log.info("  %-24s %-4s = %s", field, spec["cell"], values[field])
        return values
    finally:
        wb.close()
